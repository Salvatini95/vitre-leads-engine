"""Exportação de contatos selecionados no ProspectAdmin."""

import csv
from io import StringIO

import pytest
from django.contrib import admin as django_admin
from django.contrib.auth import get_user_model
from django.test import RequestFactory

from leads.admin import ProspectAdmin
from leads.models import Nicho, Origem, Prospect, Segmento
from leads.utils.telefone import TipoTelefone


pytestmark = pytest.mark.django_db


@pytest.fixture
def painel():
    return ProspectAdmin(
        Prospect,
        django_admin.site,
    )


@pytest.fixture
def request_admin():
    usuario = get_user_model().objects.create_superuser(
        username="admin-export",
        password="senha-teste",
    )

    request = RequestFactory().post("/admin/leads/prospect/")

    request.user = usuario

    return request


@pytest.fixture
def nicho():
    return Nicho.objects.create(
        codigo="autopecas-export",
        nome="Autopeças Export",
        ativo=True,
    )


def criar_prospect(
    nicho,
    origem_id,
    nome,
    telefone,
    telefone_tipo,
):
    return Prospect.objects.create(
        origem=Origem.FOURSQUARE,
        origem_id=origem_id,
        nome=nome,
        segmento=Segmento.OUTRO,
        nicho=nicho,
        cidade="Maringá",
        estado="PR",
        endereco="Rua Teste",
        telefone=telefone,
        telefone_tipo=telefone_tipo,
    )


def test_exportacao_e_acao_do_prospect_admin(painel):
    assert "exportar_contatos_csv" in painel.actions


def test_exporta_apenas_queryset_selecionado(
    painel,
    request_admin,
    nicho,
):
    celular = criar_prospect(
        nicho,
        "export-1",
        "Auto Peças Celular",
        "+5544999120926",
        TipoTelefone.CELULAR,
    )

    criar_prospect(
        nicho,
        "export-2",
        "Não Selecionado",
        "+554432446413",
        TipoTelefone.FIXO,
    )

    resposta = painel.exportar_contatos_csv(
        request_admin,
        Prospect.objects.filter(pk=celular.pk),
    )

    assert resposta.status_code == 200
    assert resposta["Content-Type"].startswith("text/csv")

    conteudo = resposta.content.decode("utf-8-sig")

    linhas = list(
        csv.reader(
            StringIO(conteudo)
        )
    )

    assert linhas[0] == [
        "Nome",
        "Telefone",
        "Tipo telefone",
        "Candidato a WhatsApp",
        "Nicho",
        "Cidade",
        "Estado",
        "Endereço",
        "Origem",
    ]

    assert len(linhas) == 2

    assert linhas[1][0] == "Auto Peças Celular"
    assert linhas[1][1] == "+5544999120926"
    assert linhas[1][3] == "Sim"

    assert "Não Selecionado" not in conteudo


def test_exportacao_excel_e_acao_do_prospect_admin(painel):
    assert "exportar_contatos_excel" in painel.actions


def test_exporta_excel_apenas_queryset_selecionado(
    painel,
    request_admin,
    nicho,
):
    from io import BytesIO

    from openpyxl import load_workbook

    celular = criar_prospect(
        nicho,
        "export-xlsx-1",
        "Auto Peças Excel",
        "+5544999120926",
        TipoTelefone.CELULAR,
    )

    criar_prospect(
        nicho,
        "export-xlsx-2",
        "Não Selecionado Excel",
        "+554432446413",
        TipoTelefone.FIXO,
    )

    resposta = painel.exportar_contatos_excel(
        request_admin,
        Prospect.objects.filter(pk=celular.pk),
    )

    assert resposta.status_code == 200

    assert (
        resposta["Content-Type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    workbook = load_workbook(
        BytesIO(resposta.content)
    )

    planilha = workbook["Prospects"]

    assert planilha.freeze_panes == "A2"
    assert planilha.auto_filter.ref is not None

    assert planilha["A1"].value == "Nome"
    assert planilha["B1"].value == "Telefone"
    assert planilha["D1"].value == "Candidato a WhatsApp"

    assert planilha["A2"].value == "Auto Peças Excel"
    assert planilha["B2"].value == "+5544999120926"
    assert planilha["D2"].value == "Sim"

    assert planilha.max_row == 2
